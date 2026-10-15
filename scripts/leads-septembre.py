# -*- coding: utf-8 -*-
"""Chiffres de septembre du board Lead/Prospect, a partir d'un export Monday.

    python scripts/leads-septembre.py "E:/burrea/Jim/Lead_Prospect_xxx.xlsx" [2026-09]

   MEME METHODE QUE POUR LE BOARD COMPTES
   On lit l'export brut, on repere la ligne d'en-tete et les groupes, puis on
   ne garde que les affaires dont la DATE DE RDV tombe dans le mois demande.
   C'est la cohorte : une affaire appartient au mois ou le rendez-vous a eu
   lieu, pas au mois ou elle a ete creee ni ou elle a ete signee.

   Deux precautions apprises sur le board Comptes :
   1. L'export repete parfois l'en-tete au changement de groupe, et le nom du
      groupe occupe une ligne a lui seul. On saute les deux.
   2. Un meme libelle de colonne peut apparaitre deux fois. On indexe donc par
      POSITION et non par nom, apres avoir releve la position de chaque colonne
      utile une seule fois.
"""
import io
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime

from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

ETATS_GAGNE = {'contrat signe', 'date de demarrage planifie/contrat signe',
               'resilie / arret de collaboration', 'fin de contrat',
               'suspension de contrat'}
ETATS_PERDU = {'arret suivi', 'stop contact'}


def sansacc(x):
    x = unicodedata.normalize('NFD', str(x or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', x).strip()


def cle(x):
    return sansacc(x).lower()


def nb(v):
    if v in (None, '', '.'):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).replace('\xa0', '').replace(' ', '').replace('\u20ac', '').replace(',', '.')
    t = re.sub(r'[^0-9.\-]', '', t)
    try:
        return float(t) if t not in ('', '-', '.') else 0.0
    except ValueError:
        return 0.0


def jour(v):
    if v in (None, '', '.'):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = str(v)[:10]
    for f in ('%Y-%m-%d', '%d/%m/%Y', '%d/%m/%y'):
        try:
            return datetime.strptime(t, f).date()
        except ValueError:
            pass
    return None


def lire(chemin):
    """Rend (colonnes, lignes) : une ligne est un dict colonne -> valeur."""
    wb = load_workbook(chemin, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    brut = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    entete, idx = None, None
    lignes, groupe = [], ''
    for r in brut:
        vals = [v for v in r]
        texte = [sansacc(v) for v in vals]
        pleins = [t for t in texte if t]
        if not pleins:
            continue
        # la ligne d'en-tete est celle qui porte « Nom » en premiere colonne
        if texte and texte[0] == 'Nom' and 'Closer' in texte:
            entete = texte
            continue
        if len(pleins) == 1:                 # un nom de groupe occupe sa ligne
            groupe = pleins[0]
            continue
        if entete is None:
            continue
        d = {'_vals': vals, '_groupe': groupe}
        lignes.append(d)
    return entete, lignes


def place(entete, *noms):
    """Indice de la PREMIERE colonne qui porte l'un de ces libelles.

       Indispensable ici : l'export porte deux colonnes « Nom », la societe en
       premiere position et le nom du contact en cinquieme. Un acces par
       libelle rendrait le contact, jamais la societe."""
    for n in noms:
        for i, t in enumerate(entete):
            if cle(t) == cle(n):
                return i
    for n in noms:
        for i, t in enumerate(entete):
            if t and cle(t).startswith(cle(n)):
                return i
    return None


def val(d, i):
    v = d['_vals']
    return v[i] if (i is not None and i < len(v)) else None


def main():
    if len(sys.argv) < 2:
        sys.exit('usage : python scripts/leads-septembre.py <export.xlsx> [AAAA-MM]')
    chemin = sys.argv[1]
    mois = sys.argv[2] if len(sys.argv) > 2 else '2026-09'
    an, mo = int(mois[:4]), int(mois[5:7])
    deb = date(an, mo, 1)
    fin = date(an + (mo == 12), (mo % 12) + 1, 1)

    entete, lignes = lire(chemin)
    if entete is None:
        sys.exit("en-tete introuvable : l'export ne ressemble pas au board Lead/Prospect")
    print('export : %s' % chemin)
    print('%d colonnes, %d lignes' % (len(entete), len(lignes)))
    print('colonnes : %s' % ', '.join(c for c in entete if c))
    print()

    I = {
        'nom': 0,                                   # la societe, toujours en tete
        'closer': place(entete, 'Closer'),
        'etat': place(entete, 'Etat', 'Statut com'),
        'canal': place(entete, "Canaux d'acquisition", 'Canal', 'Source'),
        'rdv': place(entete, 'Date RDV'),
        'vente': place(entete, 'Vente P', 'Prix de Vente'),
        'achat': place(entete, 'Achat P', "Prix d'Achat"),
        'proba': place(entete, '% de probabilite', '% Proba', 'Proba'),
        'motif': place(entete, 'Motif refus'),
        'demarrage': place(entete, 'Date de demarrage souhaite'),
    }
    manquantes = [k for k, v in I.items() if v is None]
    if manquantes:
        print('colonnes introuvables : %s' % ', '.join(manquantes))
    print('colonnes utilisees : %s' % ', '.join(
        '%s=%d' % (k, v) for k, v in I.items() if v is not None))
    print()

    retenues = []
    ignorees = 0
    for d in lignes:
        brut = val(d, I['rdv'])
        # Monday ajoute une ligne de TOTAUX au bas de chaque groupe : ses dates
        # sont des plages (« 2026-09-02 to 2026-09-28 »), ses montants des
        # sommes, sa probabilite une moyenne. Elle n'est pas une affaire.
        if ' to ' in str(brut) or not sansacc(val(d, I['nom'])):
            ignorees += 1
            continue
        rdv = jour(brut)
        if rdv is None or not (deb <= rdv < fin):
            continue
        retenues.append({
            'nom': sansacc(val(d, I['nom'])) or '(sans nom)',
            'closer': sansacc(val(d, I['closer'])) or '(vide)',
            'etat': sansacc(val(d, I['etat'])) or '(vide)',
            'canal': sansacc(val(d, I['canal'])) or '(non renseigne)',
            'rdv': rdv,
            'vente': nb(val(d, I['vente'])),
            'achat': nb(val(d, I['achat'])),
            'proba': nb(val(d, I['proba'])),
            'motif': sansacc(val(d, I['motif'])),
            'groupe': d.get('_groupe', ''),
        })

    print('=' * 62)
    print('AFFAIRES DONT LE RENDEZ-VOUS A EU LIEU EN %s : %d' % (mois, len(retenues)))
    print('=' * 62)
    if ignorees:
        print('(%d ligne(s) de totaux ou sans nom ecartee(s))' % ignorees)
    if not retenues:
        print('aucune affaire sur ce mois, verifier la colonne Date RDV')
        return

    par_etat = Counter(x['etat'] for x in retenues)
    print('\n-- par etat --')
    for e, n in par_etat.most_common():
        print('  %-34s %3d   %5.1f %%' % (e, n, 100 * n / len(retenues)))

    gagnes = [x for x in retenues if cle(x['etat']) in ETATS_GAGNE]
    perdus = [x for x in retenues if cle(x['etat']) in ETATS_PERDU]
    ouvertes = [x for x in retenues if cle(x['etat']) not in ETATS_GAGNE | ETATS_PERDU]
    print('\n-- lecture commerciale --')
    print('  signees  : %3d   %s' % (len(gagnes), ', '.join(x['nom'] for x in gagnes) or '-'))
    print('  perdues  : %3d   %s' % (len(perdus), ', '.join(x['nom'] for x in perdus) or '-'))
    print('  ouvertes : %3d' % len(ouvertes))

    print('\n-- par closer --')
    for c, n in Counter(x['closer'] for x in retenues).most_common():
        print('  %-24s %3d' % (c, n))

    canaux = Counter(x['canal'] for x in retenues)
    if len(canaux) > 1 or '(non renseigne)' not in canaux:
        print('\n-- par canal --')
        for c, n in canaux.most_common():
            print('  %-28s %3d' % (c, n))

    vente = sum(x['vente'] for x in retenues)
    avec = sum(1 for x in retenues if x['vente'] > 0)
    print('\n-- montants --')
    print('  pipeline annonce        %9.0f EUR sur %d affaires' % (vente, len(retenues)))
    print('  affaires avec un montant %3d sur %d' % (avec, len(retenues)))
    pond = sum(x['vente'] * x['proba'] / 100 for x in retenues)
    print('  pipeline pondere        %9.0f EUR' % pond)
    if avec == 0:
        print('  ATTENTION : aucun montant saisi, tout indicateur en euros sera vide')

    print('\n-- le detail, affaire par affaire --')
    print('  %-26s %-28s %-12s %6s %5s' % ('AFFAIRE', 'ETAT', 'CLOSER', 'VENTE', 'PROBA'))
    for x in sorted(retenues, key=lambda y: (y['etat'], y['nom'])):
        print('  %-26s %-28s %-12s %6.0f %4.0f %%' % (
            x['nom'][:26], x['etat'][:28], x['closer'][:12], x['vente'], x['proba']))

    motifs = Counter(x['motif'] for x in perdus if x['motif'])
    if motifs:
        print('\n-- motifs des affaires perdues --')
        for m, n in motifs.most_common():
            print('  %-44s %d' % (m[:44], n))


if __name__ == '__main__':
    main()
